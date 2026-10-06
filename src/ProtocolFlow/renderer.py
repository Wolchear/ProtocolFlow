from pathlib import Path

from jinja2 import Environment, PackageLoader
from weasyprint import HTML, CSS
from docx import Document
from docx.shared import Cm

from ProtocolFlow.models import Protocol, Styler
from ProtocolFlow.enums import OutputFormat

env = Environment(
    loader=PackageLoader("ProtocolFlow", "templates")
)

CSS_FILE = (
    Path(__file__).parent
    / "templates"
    / "protocol.css"
)


def render(
    protocol: Protocol,
    styler: Styler,
    output_format: OutputFormat,
    output_file_name: Path
) -> None:
    html = render_html(protocol, styler)
    output = output_file_name.with_suffix(f".{output_format.value}")
    
    match output_format:
        case OutputFormat.HTML:
            html = render_html(protocol, styler)
            output.write_text(html, encoding="utf-8")

        case OutputFormat.PDF:
            html = render_html(protocol, styler)
            HTML(string=html).write_pdf(output)

        case OutputFormat.DOCX:
            render_docx(protocol, styler, output)


def render_html(protocol: Protocol, styler: Styler) -> str:
    template = env.get_template("protocol.html")
    css = env.get_template("protocol.css").render()

    return template.render(
        protocol=protocol,
        styler=styler,
        css=css,
    )

def render_docx(
    protocol: Protocol,
    styler: Styler,
    output: Path,
) -> None:
    document = Document()

    document.add_heading(protocol.config.title, level=1)

    metadata = protocol.config.metadata

    if metadata.author:
        document.add_paragraph(f"Author: {metadata.author}")

    if metadata.sources:
        document.add_paragraph("Sources:")

        for source in metadata.sources:
            document.add_paragraph(
                source,
                style="List Bullet",
            )

    if protocol.reagents:
        document.add_heading("Reagents", level=1)

        for reagent in protocol.reagents:
            document.add_heading(reagent.name, level=2)

            if styler.reagent_style.show_recipe and reagent.recipe:
                render_recipe_docx(
                    document,
                    reagent.recipe,
                    styler.reagent_style.notes_style,
                )

            render_notes_docx(
                document,
                reagent.notes,
                styler.reagent_style.notes_style,
            )

    if protocol.stages:
        document.add_heading("Protocol", level=1)
        
        for stage in protocol.stages:
            document.add_heading(stage.name, level=2)
            render_notes_docx(
                document,
                stage.notes,
                styler.stage_style.notes_style,
            )
            for number, step in enumerate(stage.steps, start=1):
                document.add_paragraph(
                    f"{number}. {step.description}"
                )

                render_notes_docx(
                    document,
                    step.notes,
                    styler.stage_style.notes_style,
                )

    document.save(output)

def render_notes_docx(document, notes, styler):
    if not notes:
        return

    if styler.show_suggestions and notes.suggestions:
        paragraph = document.add_paragraph("Suggestions:")
        paragraph.paragraph_format.left_indent = Cm(0.75)

        for suggestion in notes.suggestions:
            paragraph = document.add_paragraph(
                suggestion,
                style="List Bullet",
            )
            paragraph.paragraph_format.left_indent = Cm(1.25)

    if styler.show_warnings and notes.warnings:
        paragraph = document.add_paragraph("Warnings:")
        paragraph.paragraph_format.left_indent = Cm(0.75)

        for warning in notes.warnings:
            paragraph = document.add_paragraph(
                warning,
                style="List Bullet",
            )
            paragraph.paragraph_format.left_indent = Cm(1.25)

def render_recipe_docx(document, recipe, notes_styler):
    document.add_paragraph(
        f"Final volume: "
        f"{recipe.final_volume.value} "
        f"{recipe.final_volume.unit}"
    )

    if recipe.reagents:
        document.add_paragraph("Reagents:")

        for component in recipe.reagents:
            text = (
                f"{component.name}: "
                f"{component.amount.value} "
                f"{component.amount.unit}"
            )

            if component.description:
                text += f" — {component.description}"

            document.add_paragraph(
                text,
                style="List Bullet",
            )

    if recipe.steps:
        document.add_paragraph("Preparation steps:")

        for number, step in enumerate(recipe.steps, start=1):
            document.add_paragraph(
                f"{number}. {step}"
            )

    render_notes_docx(
        document,
        recipe.notes,
        notes_styler,
    )
    
def add_indented_paragraph(document, text, style=None):
    paragraph = document.add_paragraph(text, style=style)
    paragraph.paragraph_format.left_indent = Cm(0.75)
    return paragraph